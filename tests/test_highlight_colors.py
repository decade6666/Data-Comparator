# -*- coding: utf-8 -*-
"""新增行必须使用配置的新增色，且不得破坏既有标签/表头/删除/更新颜色。

对应 .trellis/tasks/10-15-fix-added-row-colors 的 R1-R5：
- 标记为「新增」的数据行整行（含标记列与空单元格）使用 config.new_sheet_tab_fill；
- 「删除」行继续整行使用 config.highlight_fill；
- 「更新」行继续只高亮位置对齐的差异列，未变化单元格保持无填充；
- Sheet 标签、新增/删除列头、普通表头颜色保持现状。

单元用例直接调用真实高亮函数与真实 ConfigManager；integration 用例跑真实比对
管线（process_edc_multithreaded）并重新打开保存的报告断言序列化后的颜色。
"""

import os
import threading

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from src.backend.domain.data_comparison import process_edc_multithreaded
from src.backend.domain.excel_utils import apply_highlight_to_worksheet
from src.backend.infrastructure.config_manager import ConfigManager

MARK_COL = "更新情况（标记）"
SAS_NAMES = [MARK_COL, "colA", "colB"]

# ConfigManager._to_argb 会把 6 位 hex 规范化为 FF 前缀的 8 位 ARGB
DEFAULT_ADDITION_RGB = "FF00FF00"  # 默认新增色 #00FF00
CUSTOM_UPDATE_RGB = "FF112233"  # 自定义更新色 #112233
CUSTOM_MISSING_RGB = "FF445566"  # 自定义删除色 #445566
CUSTOM_ADDITION_RGB = "FF778899"  # 自定义新增色 #778899
HEADER_BLUE_RGB = "00A6C9EC"  # 普通表头蓝（6 位 hex 由 openpyxl 补 00 前缀）

CUSTOM_COLORS = {
    "highlight_fill": "#112233",
    "missing_sheet_tab": "#445566",
    "new_sheet_tab": "#778899",
}


def make_config(colors=None, merge_deleted_data=True):
    config = ConfigManager()
    config.update_from_parameters(
        {"default_keys": ["SUBJID"], "merge_deleted_data": merge_deleted_data},
        colors or {},
    )
    return config


def make_worksheet(rows):
    """构造一个表头 + 若干数据行的工作表，rows 为数据行取值列表。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(SAS_NAMES)
    for row in rows:
        worksheet.append(row)
    return worksheet


def _fill_rgb(cell):
    """有实心填充时返回规范化 ARGB，否则返回 None。"""
    if cell.fill is not None and cell.fill.fill_type == "solid":
        return cell.fill.start_color.rgb
    return None


# ---------------------------------------------------------------------------
# 单元测试：真实高亮函数 + 真实 ConfigManager
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("colors", "expected_rgb"),
    [
        pytest.param({}, DEFAULT_ADDITION_RGB, id="default-palette"),
        pytest.param(CUSTOM_COLORS, CUSTOM_ADDITION_RGB, id="custom-palette"),
    ],
)
def test_added_row_uses_configured_addition_fill_including_blank_cells(
    colors, expected_rgb
) -> None:
    """「新增」行整行（含空单元格）必须使用配置的新增色，而非更新色。"""
    worksheet = make_worksheet([["新增", "a0", None]])
    config = make_config(colors)

    apply_highlight_to_worksheet(ws=worksheet, config=config, sas_file_names=SAS_NAMES)

    for col_idx in range(1, len(SAS_NAMES) + 1):
        cell = worksheet.cell(row=2, column=col_idx)
        assert cell.fill.fill_type == "solid"
        assert cell.fill.start_color.rgb == expected_rgb


def test_added_row_marker_keeps_bold_red_font() -> None:
    """修复新增行填充色不得改变标记列的加粗红字样式。"""
    worksheet = make_worksheet([["新增", "a0", None]])

    apply_highlight_to_worksheet(
        ws=worksheet, config=make_config(CUSTOM_COLORS), sas_file_names=SAS_NAMES
    )

    font = worksheet.cell(row=2, column=1).font
    assert font.bold is True
    assert font.color.rgb == "FFFF0000"


def test_deleted_row_keeps_update_fill_across_all_cells() -> None:
    """「删除」行保持既有 highlight_fill 整行填充，不受新增色修复影响。"""
    worksheet = make_worksheet([["删除", "a0", None]])

    apply_highlight_to_worksheet(
        ws=worksheet, config=make_config(CUSTOM_COLORS), sas_file_names=SAS_NAMES
    )

    for col_idx in range(1, len(SAS_NAMES) + 1):
        assert _fill_rgb(worksheet.cell(row=2, column=col_idx)) == CUSTOM_UPDATE_RGB


def test_updated_row_colors_only_changed_cells() -> None:
    """「更新」行只高亮差异列，未变化列与标记列保持无填充。"""
    worksheet = make_worksheet([["更新", "a0", "b0"]])

    apply_highlight_to_worksheet(
        ws=worksheet,
        config=make_config(CUSTOM_COLORS),
        diff_info={0: {"colB": True}},
        sas_file_names=SAS_NAMES,
    )

    assert _fill_rgb(worksheet.cell(row=2, column=3)) == CUSTOM_UPDATE_RGB
    assert _fill_rgb(worksheet.cell(row=2, column=2)) is None
    assert _fill_rgb(worksheet.cell(row=2, column=1)) is None


def test_unchanged_row_has_no_cell_fill() -> None:
    """未改变行不得被染色。"""
    worksheet = make_worksheet([["未改变", "a0", "b0"]])

    apply_highlight_to_worksheet(
        ws=worksheet, config=make_config(CUSTOM_COLORS), sas_file_names=SAS_NAMES
    )

    for col_idx in range(1, len(SAS_NAMES) + 1):
        assert _fill_rgb(worksheet.cell(row=2, column=col_idx)) is None


def test_header_fills_resolve_by_sas_name_rather_than_display_label() -> None:
    """列头增删染色必须按 SAS 字段名解析，显示标签不同也照常命中。"""
    worksheet = make_worksheet([])
    worksheet.cell(row=1, column=1).value = "显示名甲"
    worksheet.cell(row=1, column=2).value = "显示名乙"
    worksheet.cell(row=1, column=3).value = MARK_COL

    apply_highlight_to_worksheet(
        ws=worksheet,
        config=make_config(CUSTOM_COLORS),
        add_sas_names=["sasB"],
        del_sas_names=["sasA"],
        sas_file_names=["sasA", "sasB", MARK_COL],
    )

    assert _fill_rgb(worksheet.cell(row=1, column=1)) == CUSTOM_MISSING_RGB
    assert _fill_rgb(worksheet.cell(row=1, column=2)) == CUSTOM_ADDITION_RGB
    assert _fill_rgb(worksheet.cell(row=1, column=3)) == HEADER_BLUE_RGB


# ---------------------------------------------------------------------------
# 集成测试：真实管线 + 重新打开保存的报告
# ---------------------------------------------------------------------------


def _make_wb(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        worksheet.append(list(frame.columns))
        for row in frame.itertuples(index=False):
            worksheet.append(list(row))
    workbook.save(path)


def _run_comparison(old_path, new_path, out_path, merge_deleted_data):
    result_path = process_edc_multithreaded(
        old_path,
        new_path,
        out_path,
        lambda message: None,
        config=make_config(CUSTOM_COLORS, merge_deleted_data),
        stop_flag=threading.Event(),
    )
    return load_workbook(result_path)


def _rows_by_subject(worksheet):
    """按 SUBJID 聚合每行的标记与逐列填充，避免依赖列顺序。"""
    header = [cell.value for cell in worksheet[1]]
    subject_idx = header.index("SUBJID") + 1
    mark_idx = header.index(MARK_COL) + 1
    rows = {}
    for row_idx in range(2, worksheet.max_row + 1):
        subject = worksheet.cell(row_idx, subject_idx).value
        rows[subject] = {
            "mark": worksheet.cell(row_idx, mark_idx).value,
            "fills": {
                header[col_idx - 1]: _fill_rgb(worksheet.cell(row_idx, col_idx))
                for col_idx in range(1, worksheet.max_column + 1)
            },
        }
    return rows


# 001 只在旧文件（删除）；002 改 A；003 不变；004 只在新文件（新增，B 为空）。
_OLD = {"SUBJID": ["001", "002", "003"], "A": ["x", "y", "z"], "B": ["p", "q", "r"]}
_NEW = {"SUBJID": ["002", "003", "004"], "A": ["y2", "z", "w"], "B": ["q", "r", None]}


@pytest.mark.integration
@pytest.mark.parametrize("merge_deleted_data", [True, False])
def test_common_sheet_mixed_records_use_expected_colors(
    tmp_path, merge_deleted_data
) -> None:
    """共有表单中新增行用新增色；删除/更新/未改变行与关闭合并时的对齐不受影响。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")
    _make_wb(old_path, {"MH": pd.DataFrame(_OLD)})
    _make_wb(new_path, {"MH": pd.DataFrame(_NEW)})

    workbook = _run_comparison(old_path, new_path, out_path, merge_deleted_data)
    try:
        rows = _rows_by_subject(workbook["MH"])

        if merge_deleted_data:
            assert set(rows) == {"001", "002", "003", "004"}
            # 删除行保持既有更新色整行填充
            assert rows["001"]["mark"] == "删除"
            assert set(rows["001"]["fills"].values()) == {CUSTOM_UPDATE_RGB}
        else:
            assert set(rows) == {"002", "003", "004"}

        # 新增行整行（含空单元格）使用自定义新增色
        assert rows["004"]["mark"] == "新增"
        assert set(rows["004"]["fills"].values()) == {CUSTOM_ADDITION_RGB}

        # 更新行只高亮变化列
        assert rows["002"]["mark"] == "更新"
        filled = {
            name: rgb for name, rgb in rows["002"]["fills"].items() if rgb is not None
        }
        assert filled == {"A": CUSTOM_UPDATE_RGB}

        # 未改变行不染色
        assert rows["003"]["mark"] == "未改变"
        assert set(rows["003"]["fills"].values()) == {None}
    finally:
        workbook.close()


@pytest.mark.integration
def test_new_only_sheet_tab_rows_and_headers(tmp_path) -> None:
    """新增表单：标签与数据行用新增色，普通表头保持既有蓝色。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")
    _make_wb(old_path, {"MH": pd.DataFrame(_OLD)})
    _make_wb(
        new_path,
        {
            "MH": pd.DataFrame(_NEW),
            "NEWX": pd.DataFrame({"SUBJID": ["N1", "N2"], "NOTE": ["n1", "n2"]}),
        },
    )

    workbook = _run_comparison(old_path, new_path, out_path, True)
    try:
        assert "NEWX" in workbook.sheetnames
        worksheet = workbook["NEWX"]

        tab_color = worksheet.sheet_properties.tabColor
        assert tab_color is not None
        assert tab_color.rgb == CUSTOM_ADDITION_RGB

        header = [cell.value for cell in worksheet[1]]
        mark_idx = header.index(MARK_COL) + 1
        for row_idx in range(2, worksheet.max_row + 1):
            assert worksheet.cell(row_idx, mark_idx).value == "新增"
            for col_idx in range(1, worksheet.max_column + 1):
                assert (
                    _fill_rgb(worksheet.cell(row_idx, col_idx)) == CUSTOM_ADDITION_RGB
                )

        for col_idx in range(1, worksheet.max_column + 1):
            assert _fill_rgb(worksheet.cell(row=1, column=col_idx)) == HEADER_BLUE_RGB
    finally:
        workbook.close()


@pytest.mark.integration
def test_empty_to_populated_sheet_rows_use_addition_fill(tmp_path) -> None:
    """旧表单仅有表头、新版有数据：数据行同样走新增色。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")
    _make_wb(old_path, {"MH": pd.DataFrame({"SUBJID": [], "A": [], "B": []})})
    _make_wb(
        new_path,
        {
            "MH": pd.DataFrame(
                {"SUBJID": ["001", "002"], "A": ["x", "y"], "B": ["p", "q"]}
            )
        },
    )

    workbook = _run_comparison(old_path, new_path, out_path, True)
    try:
        worksheet = workbook["MH"]

        tab_color = worksheet.sheet_properties.tabColor
        assert tab_color is not None
        assert tab_color.rgb == CUSTOM_ADDITION_RGB

        header = [cell.value for cell in worksheet[1]]
        mark_idx = header.index(MARK_COL) + 1
        assert worksheet.max_row >= 3
        for row_idx in range(2, worksheet.max_row + 1):
            assert worksheet.cell(row_idx, mark_idx).value == "新增"
            for col_idx in range(1, worksheet.max_column + 1):
                assert (
                    _fill_rgb(worksheet.cell(row_idx, col_idx)) == CUSTOM_ADDITION_RGB
                )
    finally:
        workbook.close()


@pytest.mark.integration
def test_added_column_header_colored_but_body_not_repainted(tmp_path) -> None:
    """新增列头保持新增色；列正文不得被整列重涂为新增色。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")
    _make_wb(
        old_path,
        {"MH": pd.DataFrame({"SUBJID": ["001", "002", "003"], "A": ["x", "y", "z"]})},
    )
    _make_wb(
        new_path,
        {
            "MH": pd.DataFrame(
                {
                    "SUBJID": ["001", "002", "003"],
                    "A": ["x2", "y", "z"],
                    "C": ["c1", "", "c3"],
                }
            )
        },
    )

    workbook = _run_comparison(old_path, new_path, out_path, True)
    try:
        worksheet = workbook["MH"]
        header = [cell.value for cell in worksheet[1]]
        added_idx = header.index("C") + 1
        ordinary_idx = header.index("A") + 1

        # 新增列头使用新增色，普通列头保持蓝色
        assert _fill_rgb(worksheet.cell(row=1, column=added_idx)) == CUSTOM_ADDITION_RGB
        assert _fill_rgb(worksheet.cell(row=1, column=ordinary_idx)) == HEADER_BLUE_RGB

        # 列正文任何单元格都不得被新增色整列重涂
        for row_idx in range(2, worksheet.max_row + 1):
            assert (
                _fill_rgb(worksheet.cell(row_idx, column=added_idx))
                != CUSTOM_ADDITION_RGB
            )
    finally:
        workbook.close()
