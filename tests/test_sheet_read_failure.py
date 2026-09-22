"""读失败语义回归测试：读取失败必须显式暴露，不得伪装成「缺失表单」。

背景：read_single_sheet_from_excel 读取失败时曾返回 None，与「Sheet 不存在」
同义，导致 process_single_sheet_complete 走缺失分支、旧表整表被标记「删除」、
任务终态却是成功。现在读取失败抛 SheetReadError，聚合层需收集失败表单并在
保存前打醒目汇总（报告结构不变，失败表单在报告中直接缺席）。
"""

import os
import threading

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from src.backend.domain import data_comparison
from src.backend.domain.data_comparison import (
    process_edc_multithreaded,
    process_single_sheet_complete,
)
from src.backend.domain.excel_header_utils import (
    SheetReadError,
    read_single_sheet_from_excel,
)
from src.backend.infrastructure.config_manager import ConfigManager
from src.backend.infrastructure.progress_manager import ThreadSafeProgressManager


def _make_wb(path, sheets):
    """按 {sheet名: DataFrame} 写一个真实工作簿（表头行 + 数据行）。"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append(list(row))
    wb.save(path)


def _config(params=None):
    cm = ConfigManager()
    merged = dict({"default_keys": ["SUBJID"]}, **(params or {}))
    cm.update_from_parameters(merged, {})
    return cm


def test_missing_sheet_still_returns_none(tmp_path):
    """请求不存在的 Sheet 名必须继续返回 None（三态语义的「不存在」一极不变）。"""
    path = os.path.join(str(tmp_path), "one.xlsx")
    _make_wb(path, {"S1": pd.DataFrame({"SUBJID": ["001"]})})

    df = read_single_sheet_from_excel(
        str(path), "不存在的表单", 1, 1, lambda _message: None
    )

    assert df is None


@pytest.mark.parametrize("fail_side", ["new", "old"])
def test_read_failure_does_not_become_missing_sheet(tmp_path, monkeypatch, fail_side):
    """任一侧读取失败时不得进入缺失/新增分支，而是失败结果（R2 核心）。

    新文件侧失败若被吞成 None 会走「缺失表单」分支（旧表整表标删除）；
    旧文件侧失败会走「新增表单」分支。两个方向都必须被 SheetReadError 拦下。
    """
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    _make_wb(old_path, {"AE": pd.DataFrame({"SUBJID": ["001"], "AETERM": ["a"]})})
    _make_wb(new_path, {"AE": pd.DataFrame({"SUBJID": ["001"], "AETERM": ["b"]})})

    fail_path = new_path if fail_side == "new" else old_path

    def fake_read(file_path, sheet_name, *_args, **_kwargs):
        if file_path == fail_path:
            raise SheetReadError(f"读取Sheet [{sheet_name}] 失败: 模拟读取异常")
        return pd.DataFrame({"SUBJID": ["001"], "AETERM": ["a"]})

    monkeypatch.setattr(data_comparison, "read_single_sheet_from_excel", fake_read)

    logs = []
    progress_manager = ThreadSafeProgressManager(1, None, logs.append)
    result = process_single_sheet_complete(
        "AE", old_path, new_path, _config(), progress_manager
    )

    assert result.success is False
    assert result.error_message is not None
    assert "AE" in result.error_message
    assert "读取Sheet [AE] 失败" in result.error_message
    assert result.change_type not in ("missing", "new")
    assert not any("检测到缺失Sheet" in message for message in logs)
    assert not any("检测到新增Sheet" in message for message in logs)
    assert any("处理Sheet [AE] 时出错" in message for message in logs)


@pytest.mark.integration
def test_failed_sheets_logged_in_summary(tmp_path, monkeypatch):
    """含失败表单的聚合流程：保存前必须有集中失败汇总，报告结构不变（R3）。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")
    _make_wb(
        old_path,
        {
            "AE": pd.DataFrame({"SUBJID": ["001"], "AEMODIFY": ["a"]}),
            "DM": pd.DataFrame({"SUBJID": ["001"], "AGE": [30]}),
        },
    )
    _make_wb(
        new_path,
        {
            "AE": pd.DataFrame({"SUBJID": ["001"], "AEMODIFY": ["b"]}),
            "DM": pd.DataFrame({"SUBJID": ["001"], "AGE": [30]}),
        },
    )

    real_read = data_comparison.read_single_sheet_from_excel

    def fake_read(file_path, sheet_name, *args, **kwargs):
        if sheet_name == "AE":
            raise SheetReadError(f"读取Sheet [{sheet_name}] 失败: 模拟读取异常")
        return real_read(file_path, sheet_name, *args, **kwargs)

    monkeypatch.setattr(data_comparison, "read_single_sheet_from_excel", fake_read)

    logs = []
    result = process_edc_multithreaded(
        old_path,
        new_path,
        out_path,
        logs.append,
        config=_config(),
        stop_flag=threading.Event(),
    )

    assert result == out_path
    # 保存前的集中失败汇总与逐条原因
    assert any("⚠️ 共 1 个表单处理失败:" in message for message in logs)
    assert any(
        message.strip().startswith("AE: ") and "读取Sheet [AE] 失败" in message
        for message in logs
    )
    # 失败表单不得伪装成缺失表单
    assert not any("检测到缺失Sheet" in message for message in logs)
    # 报告结构不变：失败表单直接缺席（不被误标「删除」），正常表单保留，
    # 不新增任何失败相关的工作表
    wb = load_workbook(out_path)
    assert "AE" not in wb.sheetnames
    assert wb.sheetnames == ["比对结果汇总", "DM"]


@pytest.mark.integration
def test_report_structure_unchanged_when_no_failure(tmp_path):
    """全部表单正常处理时：输出工作表集合与改动前一致，且无失败汇总日志。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")
    _make_wb(
        old_path,
        {
            "AE": pd.DataFrame({"SUBJID": ["001"], "AEMODIFY": ["a"]}),
            "DM": pd.DataFrame({"SUBJID": ["001"], "AGE": [30]}),
        },
    )
    _make_wb(
        new_path,
        {
            "AE": pd.DataFrame({"SUBJID": ["001"], "AEMODIFY": ["b"]}),
            "DM": pd.DataFrame({"SUBJID": ["001"], "AGE": [30]}),
        },
    )

    logs = []
    result = process_edc_multithreaded(
        old_path,
        new_path,
        out_path,
        logs.append,
        config=_config(),
        stop_flag=threading.Event(),
    )

    assert result == out_path
    wb = load_workbook(out_path)
    # 与既有行为一致：汇总表 + 按新文件顺序的数据表，无任何新增工作表
    assert wb.sheetnames == ["比对结果汇总", "AE", "DM"]
    assert not any("个表单处理失败" in message for message in logs)


@pytest.mark.integration
def test_real_read_failure_end_to_end_without_monkeypatch(tmp_path):
    """用真实失败触发器（锚点行重复列名）跑完整条链路，不注入任何 mock。

    其余用例都靠 monkeypatch 注入 SheetReadError，锁的是「给定读取函数抛异常时
    上层怎么做」；这条补上另一半——真实的读取失败确实会抛 SheetReadError 而不是
    返回 None，并一路走到失败汇总。覆盖 read → raise → success=False → ⚠️ 汇总
    → 报告中缺席 的完整闭环。
    """
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")

    def _make_wb_with_dup_anchor(path):
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet(title="AE")
        ws.append(["SUBJID", "AETERM", "SUBJID"])  # 锚点行重复列名 → 真实读取失败
        ws.append(["001", "a", "x"])
        ws2 = wb.create_sheet(title="DM")
        ws2.append(["SUBJID", "AGE"])
        ws2.append(["001", 30])
        wb.save(path)

    _make_wb_with_dup_anchor(old_path)
    _make_wb_with_dup_anchor(new_path)

    logs = []
    result = process_edc_multithreaded(
        old_path,
        new_path,
        out_path,
        logs.append,
        config=_config(),
        stop_flag=threading.Event(),
    )

    assert result == out_path
    # 真实失败被收集进汇总，且原因保留了具体诊断信息
    assert any("⚠️ 共 1 个表单处理失败:" in message for message in logs)
    assert any("锚点行(第 1 行)存在重复内容" in message for message in logs)
    # 不得伪装成缺失表单
    assert not any("检测到缺失Sheet" in message for message in logs)
    # 失败表单在报告中缺席，正常表单不受影响，未新增任何工作表
    wb = load_workbook(out_path)
    assert wb.sheetnames == ["比对结果汇总", "DM"]
