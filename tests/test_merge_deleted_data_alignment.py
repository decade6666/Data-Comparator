# -*- coding: utf-8 -*-
"""关闭「合并删除数据」时的高亮行对齐测试。

背景：`perform_full_comparison` 在 `merge_deleted_data=False` 时会过滤掉「删除」行，
但过滤后没有重置索引，幸存行的索引因此出现空洞。写 Excel 时按位置顺序写，
`apply_highlight_to_worksheet` 又按位置反查（`data_row_idx = row_idx - 2`），
而 `diff_dict` 的键是过滤前的索引——从第一条删除记录往后，每一行都会取到别人的差异集，
把没变的单元格标成变化、把真正变化的单元格漏掉。
"""

import os
import threading

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.backend.domain.data_comparison import (
    perform_full_comparison,
    process_edc_multithreaded,
)
from src.backend.infrastructure.config_manager import ConfigManager


class _Recorder:
    def __init__(self):
        self.messages = []

    def safe_log(self, message):
        self.messages.append(message)


def _df_with_sas(data):
    frame = pd.DataFrame(data)
    frame.attrs["sas_file_name"] = list(frame.columns)
    frame.attrs["sas_name_to_label"] = {name: name for name in frame.columns}
    return frame


# 001 只存在于旧文件（删除）；002 只改 A；003 只改 B。
# 关闭合并后 001 被过滤掉，002/003 的位置比它们的索引各小 1。
_OLD = {"SUBJID": ["001", "002", "003"], "A": ["x", "y", "z"], "B": ["p", "q", "r"]}
_NEW = {"SUBJID": ["002", "003"], "A": ["y2", "z"], "B": ["q", "r2"]}


def _config(merge_deleted_data):
    config = ConfigManager()
    config.update_from_parameters(
        {"default_keys": ["SUBJID"], "merge_deleted_data": merge_deleted_data}, {}
    )
    return config


def test_output_index_is_contiguous_after_dropping_deleted_rows() -> None:
    """过滤删除行后索引必须仍然连续——高亮按位置反查，空洞即错位。"""
    result_df, diffs = perform_full_comparison(
        "MH",
        _df_with_sas(_OLD),
        _df_with_sas(_NEW),
        ["SUBJID"],
        _config(False),
        _Recorder(),
    )[:2]

    assert list(result_df.index) == list(range(len(result_df)))
    # 每条差异记录都必须能按位置找到对应行
    assert all(0 <= key < len(result_df) for key in diffs)


def test_diff_dict_keys_point_at_the_right_rows() -> None:
    """按位置取到的差异集，必须属于该位置上的那条记录。"""
    result_df, diffs = perform_full_comparison(
        "MH",
        _df_with_sas(_OLD),
        _df_with_sas(_NEW),
        ["SUBJID"],
        _config(False),
        _Recorder(),
    )[:2]

    by_subject = {
        row["SUBJID"]: set(diffs.get(position, {}))
        for position, (_, row) in enumerate(result_df.iterrows())
    }

    assert by_subject["002"] == {"A"}
    assert by_subject["003"] == {"B"}


def _make_wb(path, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        worksheet.append(list(frame.columns))
        for row in frame.itertuples(index=False):
            worksheet.append(list(row))
    workbook.save(path)


def test_highlight_lands_on_the_changed_cell_without_merge(tmp_path) -> None:
    """端到端：关闭合并删除数据后，高亮仍必须落在真正变化的单元格上。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")

    _make_wb(old_path, {"MH": pd.DataFrame(_OLD)})
    _make_wb(new_path, {"MH": pd.DataFrame(_NEW)})

    result_path = process_edc_multithreaded(
        old_path,
        new_path,
        out_path,
        lambda message: None,
        config=_config(False),
        stop_flag=threading.Event(),
    )

    worksheet = load_workbook(result_path)["MH"]
    header = [cell.value for cell in worksheet[1]]
    subject_idx = header.index("SUBJID") + 1

    highlighted = {}
    for row_idx in range(2, worksheet.max_row + 1):
        subject = worksheet.cell(row_idx, subject_idx).value
        highlighted[subject] = {
            header[col_idx - 1]
            for col_idx in range(1, worksheet.max_column + 1)
            if worksheet.cell(row_idx, col_idx).fill.fill_type
        }

    # 删除行已被剔除，两条更新行各自只高亮自己变化的那一列
    assert set(highlighted) == {"002", "003"}
    assert highlighted["002"] == {"A"}
    assert highlighted["003"] == {"B"}


def test_merge_enabled_path_is_unchanged(tmp_path) -> None:
    """开启合并删除数据时不做过滤，原有行为必须保持不变。"""
    old_path = os.path.join(str(tmp_path), "old.xlsx")
    new_path = os.path.join(str(tmp_path), "new.xlsx")
    out_path = os.path.join(str(tmp_path), "out.xlsx")

    _make_wb(old_path, {"MH": pd.DataFrame(_OLD)})
    _make_wb(new_path, {"MH": pd.DataFrame(_NEW)})

    result_path = process_edc_multithreaded(
        old_path,
        new_path,
        out_path,
        lambda message: None,
        config=_config(True),
        stop_flag=threading.Event(),
    )

    worksheet = load_workbook(result_path)["MH"]
    header = [cell.value for cell in worksheet[1]]
    subject_idx = header.index("SUBJID") + 1
    mark_idx = header.index("更新情况（标记）") + 1

    marks = {
        worksheet.cell(row_idx, subject_idx)
        .value: worksheet.cell(row_idx, mark_idx)
        .value
        for row_idx in range(2, worksheet.max_row + 1)
    }

    assert marks == {"001": "删除", "002": "更新", "003": "更新"}
