"""回归测试：read_single_sheet_from_excel 不应被错误的 <dimension> 声明截断。

背景：CRF-Editor 导出的 xlsx 每个工作表都写死 <dimension ref="A1"/>，
但实际有多行多列数据。openpyxl 只读模式采信 dimension 会把表截断成 1x1，
导致读回的 DataFrame 为空、比较阶段打印「在新旧版本中均为空」。
"""

import re
import shutil
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from src.backend.domain.excel_header_utils import read_single_sheet_from_excel


def _make_workbook(path, sheet_name: str, rows, header_row: int = 1) -> None:
    """写一个简单工作簿：header 行在前，随后是数据行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def _force_dimension_a1(path) -> None:
    """把 xl/worksheets/sheet1.xml 的 dimension 声明改成 ref="A1"。"""
    tmp = path + ".rewrite"
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        infos = {info.filename: info for info in archive.infolist()}
        with zipfile.ZipFile(tmp, "w", allowZip64=True) as target:
            for name in names:
                data = archive.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    data = re.sub(
                        rb'<dimension ref="[^"]*"\s*/>',
                        b'<dimension ref="A1"/>',
                        data,
                    )
                target.writestr(infos[name], data)
    shutil.move(tmp, path)


@pytest.fixture()
def dim_broken_xlsx(tmp_path):
    """维度声明 A1、实际 5 行 × 4 列的工作簿。"""
    path = tmp_path / "dim_broken.xlsx"
    rows = [
        ["Label1", "Label2", "Label3", "Label4"],
        ["K1", "K2", "K3", "K4"],
        ["a", "b", "c", "d"],
        ["e", "f", "g", "h"],
        ["i", "j", "k", "l"],
    ]
    _make_workbook(str(path), "S1", rows)
    _force_dimension_a1(str(path))
    return str(path)


def test_reads_full_data_when_dimension_declares_a1(dim_broken_xlsx) -> None:
    """dimension 声明 A1 时仍应读到全部行与列。"""
    messages = []
    df = read_single_sheet_from_excel(
        dim_broken_xlsx,
        "S1",
        anchor_row_num=2,
        header_row_num=1,
        log_func=messages.append,
    )
    assert df is not None
    # 5 行 - 2（表头 + 锚点）= 3 行数据；4 列
    assert df.shape == (3, 4)
    assert list(df.columns) == ["K1", "K2", "K3", "K4"]
    assert df.iloc[0].tolist() == ["a", "b", "c", "d"]
    assert df.attrs["sas_file_name"] == ["K1", "K2", "K3", "K4"]
    assert df.attrs["sas_file_label"] == ["Label1", "Label2", "Label3", "Label4"]


def test_normal_dimension_reads_unchanged(tmp_path) -> None:
    """维度声明正确的普通文件，读取结果保持正确。"""
    path = tmp_path / "normal.xlsx"
    rows = [
        ["Label1", "Label2"],
        ["K1", "K2"],
        ["a", "b"],
    ]
    _make_workbook(str(path), "S1", rows)
    df = read_single_sheet_from_excel(
        str(path), "S1", anchor_row_num=2, header_row_num=1, log_func=lambda m: None
    )
    assert df is not None
    assert df.shape == (1, 2)
    assert df.iloc[0].tolist() == ["a", "b"]


def test_truly_empty_sheet_stays_empty(tmp_path) -> None:
    """只有表头/锚点行、没有数据行的空表仍返回空 DataFrame。"""
    path = tmp_path / "empty.xlsx"
    rows = [
        ["Label1", "Label2"],
        ["K1", "K2"],
    ]
    _make_workbook(str(path), "S1", rows)
    df = read_single_sheet_from_excel(
        str(path), "S1", anchor_row_num=2, header_row_num=1, log_func=lambda m: None
    )
    assert df is not None
    assert df.empty
    assert list(df.columns) == ["K1", "K2"]


def test_ragged_rows_are_padded_after_dimension_reset(tmp_path) -> None:
    """重置维度后行宽不齐时，短行缺失列以 None/NaN 填充，列数取最大行宽。"""
    path = tmp_path / "ragged.xlsx"
    _make_workbook(
        str(path),
        "S1",
        [
            ["Label1", "Label2", "Label3"],
            ["K1", "K2", "K3"],
            ["a", "b", "c"],
            ["x", "y"],  # 短一行
        ],
    )
    _force_dimension_a1(str(path))
    df = read_single_sheet_from_excel(
        str(path), "S1", anchor_row_num=2, header_row_num=1, log_func=lambda m: None
    )
    assert df is not None
    assert df.shape == (2, 3)
    # 短行缺失的 K3 以 None/NaN 填充，不抛错
    assert pd.isna(df.iloc[1]["K3"]) or df.iloc[1]["K3"] == ""


def test_styled_empty_tail_in_anchor_row_does_not_break_read(tmp_path) -> None:
    """锚点/表头行尾部只有样式没有值的单元格时，读取不应失败（dimension 声明正确也一样）。"""
    path = tmp_path / "styled_tail.xlsx"
    _make_workbook(
        str(path),
        "S1",
        [
            ["Label1", "Label2", "Label3"],
            ["K1", "K2", "K3"],
            ["a", "b", "c"],
            ["d", "e", "f"],
        ],
    )
    fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    wb = load_workbook(str(path))
    ws = wb["S1"]
    for row_idx in (1, 2):
        for col_idx in (4, 5, 6):  # D/E/F 加底色但不填值
            ws.cell(row=row_idx, column=col_idx).fill = fill
    wb.save(str(path))
    wb.close()

    df = read_single_sheet_from_excel(
        str(path), "S1", anchor_row_num=2, header_row_num=1, log_func=lambda m: None
    )
    assert df is not None
    assert df.shape == (2, 3)
    assert list(df.columns) == ["K1", "K2", "K3"]


def test_declared_but_dataless_tail_fields_are_kept(tmp_path) -> None:
    """锚点行尾部真实声明、但没有对应数据列的字段必须保留（已声明 schema 不丢失）。"""
    path = tmp_path / "declared_tail.xlsx"
    _make_workbook(
        str(path),
        "S1",
        [
            ["Label1", "Label2", "Label3", "Label4", "Label5"],
            ["K1", "K2", "K3", "K4", "K5"],
            ["a", "b", "c"],
            ["d", "e", "f"],
        ],
    )
    df = read_single_sheet_from_excel(
        str(path), "S1", anchor_row_num=2, header_row_num=1, log_func=lambda m: None
    )
    assert df is not None
    assert df.shape == (2, 5)
    assert list(df.columns) == ["K1", "K2", "K3", "K4", "K5"]
    # 无数据的声明字段对应列全为 NaN/None
    assert df[["K4", "K5"]].isna().all().all()
    assert df.attrs["sas_file_name"] == ["K1", "K2", "K3", "K4", "K5"]


def test_no_empty_or_duplicate_column_names(tmp_path) -> None:
    """锚点行中段空洞不得产生空字符串或重复列名（下游按列名单列访问会崩溃）。"""
    path = tmp_path / "holes.xlsx"
    _make_workbook(
        str(path),
        "S1",
        [
            ["Label1", "Label2", "Label3", "Label4", "Label5"],
            ["K1", None, "K3", None, "K5"],
            ["a", "b", "c", "d", "e"],
        ],
    )
    df = read_single_sheet_from_excel(
        str(path), "S1", anchor_row_num=2, header_row_num=1, log_func=lambda m: None
    )
    assert df is not None
    columns = list(df.columns)
    assert "" not in columns
    assert len(set(columns)) == len(columns)
    # 空洞位按绝对下标补名
    assert columns[1] == "Unnamed_1"
    assert columns[3] == "Unnamed_3"


def test_label_and_name_lengths_stay_aligned(tmp_path) -> None:
    """表头行、锚点行、数据行宽度互不一致时，attrs 中 label/name 长度必须与列数对齐。"""
    path = tmp_path / "misaligned.xlsx"
    _make_workbook(
        str(path),
        "S1",
        [
            ["Label1", "Label2", "Label3", "Label4", "Label5"],
            ["K1", "K2", "K3", "K4"],
            ["a", "b", "c"],
        ],
    )
    df = read_single_sheet_from_excel(
        str(path), "S1", anchor_row_num=2, header_row_num=1, log_func=lambda m: None
    )
    assert df is not None
    assert len(df.attrs["sas_file_label"]) == len(df.attrs["sas_file_name"])
    assert len(df.attrs["sas_file_name"]) == len(df.columns)
