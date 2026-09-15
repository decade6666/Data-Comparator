"""高亮写回的停止响应、性能等价性与向后兼容测试。

对应 .trellis/tasks/09-15-anchor-explosion 的 AC4 / AC5：
- diff_keys 只构建一次，且高亮结果与逐行重建时逐单元格一致；
- apply_highlight_to_worksheet 支持 stop_flag，置位时抛 InterruptedError；
- 不传 stop_flag 时行为不变。
"""

import threading
import types

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from src.backend.domain.excel_utils import apply_highlight_to_worksheet

MARK_COL = "更新情况（标记）"
SAS_NAMES = [MARK_COL, "colA", "colB"]

HIGHLIGHT = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
MISSING = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
NEW = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")


def make_config():
    return types.SimpleNamespace(
        highlight_fill=HIGHLIGHT,
        missing_sheet_tab_fill=MISSING,
        new_sheet_tab_fill=NEW,
    )


def make_worksheet(marks):
    """构造一个表头 + len(marks) 个数据行的工作表。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(SAS_NAMES)
    for index, mark in enumerate(marks):
        worksheet.append([mark, f"a{index}", f"b{index}"])
    return worksheet


def snapshot_fills(worksheet):
    """逐单元格记录填充色与字体加粗，用于等价性比较。"""
    return [
        (cell.row, cell.column, cell.fill.start_color.rgb, cell.font.bold)
        for row in worksheet.iter_rows()
        for cell in row
    ]


def test_apply_highlight_raises_on_stop_flag() -> None:
    """停止标志置位时，高亮阶段必须抛出 InterruptedError 而不是跑完整个表单。"""
    worksheet = make_worksheet(["更新"] * 200)
    stop_flag = threading.Event()
    stop_flag.set()

    with pytest.raises(InterruptedError, match="用户停止了操作"):
        apply_highlight_to_worksheet(
            ws=worksheet,
            config=make_config(),
            sheet_type="data_changed",
            diff_info={index: {"colA": True} for index in range(200)},
            sas_file_names=SAS_NAMES,
            stop_flag=stop_flag,
        )


def test_apply_highlight_unset_stop_flag_completes() -> None:
    """未置位的 stop_flag 不应影响正常执行。"""
    worksheet = make_worksheet(["更新", "新增"])

    apply_highlight_to_worksheet(
        ws=worksheet,
        config=make_config(),
        sheet_type="data_changed",
        diff_info={0: {"colA": True}},
        sas_file_names=SAS_NAMES,
        stop_flag=threading.Event(),
    )

    assert (
        worksheet.cell(row=2, column=2).fill.start_color.rgb
        == HIGHLIGHT.start_color.rgb
    )


def test_apply_highlight_without_stop_flag_is_backward_compatible() -> None:
    """既有调用方不传 stop_flag，签名与行为必须保持兼容。"""
    worksheet = make_worksheet(["更新", "删除"])

    apply_highlight_to_worksheet(
        ws=worksheet,
        config=make_config(),
        sheet_type="data_changed",
        diff_info={0: {"colB": True}},
        sas_file_names=SAS_NAMES,
    )

    # 「更新」行只高亮有差异的列（colB 是第 3 列）
    assert (
        worksheet.cell(row=2, column=3).fill.start_color.rgb
        == HIGHLIGHT.start_color.rgb
    )
    assert (
        worksheet.cell(row=2, column=2).fill.start_color.rgb
        != HIGHLIGHT.start_color.rgb
    )
    # 「删除」行整行高亮
    for column in (1, 2, 3):
        assert (
            worksheet.cell(row=3, column=column).fill.start_color.rgb
            == HIGHLIGHT.start_color.rgb
        )


def test_apply_highlight_result_equivalence_across_diff_key_types() -> None:
    """diff_info 的键无论是 int 还是 str 数字，高亮结果都必须一致。

    这锁定「diff_keys 提取到循环外」这一改动的等价性：提取前后逐单元格结果相同。
    """
    marks = ["更新", "未改变", "新增", "更新", "删除"]
    diff_int = {0: {"colA": True}, 3: {"colB": True}}
    diff_str = {"0": {"colA": True}, "3": {"colB": True}}

    worksheet_int = make_worksheet(marks)
    apply_highlight_to_worksheet(
        ws=worksheet_int,
        config=make_config(),
        sheet_type="data_changed",
        diff_info=diff_int,
        sas_file_names=SAS_NAMES,
    )

    worksheet_str = make_worksheet(marks)
    apply_highlight_to_worksheet(
        ws=worksheet_str,
        config=make_config(),
        sheet_type="data_changed",
        diff_info=diff_str,
        sas_file_names=SAS_NAMES,
    )

    assert snapshot_fills(worksheet_int) == snapshot_fills(worksheet_str)


def test_apply_highlight_builds_diff_keys_once() -> None:
    """diff_keys 必须只构建一次，而不是每个数据行重建一遍。

    用一个会统计 items() 调用次数的 dict 子类探测：O(N×M) 的旧实现会调用 N 次。
    """

    class CountingDict(dict):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.items_calls = 0

        def items(self):
            self.items_calls += 1
            return super().items()

    row_count = 50
    diff_info = CountingDict({index: {"colA": True} for index in range(row_count)})
    worksheet = make_worksheet(["更新"] * row_count)

    apply_highlight_to_worksheet(
        ws=worksheet,
        config=make_config(),
        sheet_type="data_changed",
        diff_info=diff_info,
        sas_file_names=SAS_NAMES,
    )

    assert diff_info.items_calls == 1
